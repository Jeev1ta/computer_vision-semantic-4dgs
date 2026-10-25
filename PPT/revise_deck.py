from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SOURCE = Path(__file__).with_name("cv final ppt.pptx")
OUTPUT = Path(__file__).with_name("cv final ppt_revised.pptx")
NAVY = RGBColor(23, 51, 84)
TEAL = RGBColor(34, 124, 152)
DARK = RGBColor(45, 52, 60)
MID = RGBColor(90, 99, 108)
PALE = RGBColor(238, 242, 245)
WHITE = RGBColor(255, 255, 255)
GOLD_PALE = RGBColor(255, 248, 225)


def set_cell(cell, text, *, header=False, shade=False, size=12):
    cell.text = text
    cell.margin_left = Inches(0.10)
    cell.margin_right = Inches(0.10)
    cell.margin_top = Inches(0.035)
    cell.margin_bottom = Inches(0.025)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = NAVY if header else (PALE if shade else WHITE)
    for paragraph in cell.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = header
            run.font.color.rgb = WHITE if header else DARK


def clear_slide_shapes(slide, keep_shape_ids=()):
    keep_shape_ids = set(keep_shape_ids)
    for shape in list(slide.shapes):
        if shape.shape_id not in keep_shape_ids:
            shape._element.getparent().remove(shape._element)


def add_title(slide, title, subtitle, page_num):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.60), Inches(0.50), Inches(0.09), Inches(0.95))
    bar.fill.solid(); bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()
    for text, y, h, size, color, bold in (
        (title, 0.45, 0.60, 25, NAVY, True),
        (subtitle, 1.05, 0.42, 15, MID, False),
    ):
        box = slide.shapes.add_textbox(Inches(0.85), Inches(y), Inches(11.8), Inches(h))
        tf = box.text_frame; tf.clear(); tf.word_wrap = True
        p = tf.paragraphs[0]; p.space_after = Pt(0)
        r = p.add_run(); r.text = text
        r.font.name = "Arial"; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    add_footer(slide, page_num)


def add_footer(slide, page_num):
    box = slide.shapes.add_textbox(Inches(0.60), Inches(7.0), Inches(10.5), Inches(0.30))
    p = box.text_frame.paragraphs[0]
    r = p.add_run(); r.text = "Semantic-Aware Editable Dynamic 3D Reconstruction via Gaussian Splatting   |   Course Project Final Review"
    r.font.name = "Arial"; r.font.size = Pt(8); r.font.color.rgb = MID
    number = slide.shapes.add_textbox(Inches(12.1), Inches(7.0), Inches(0.6), Inches(0.30))
    p = number.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = str(page_num); r.font.name = "Arial"; r.font.size = Pt(8); r.font.color.rgb = MID


def add_table(slide, x, y, w, h, rows, col_widths, font_size=12):
    table_shape = slide.shapes.add_table(len(rows), len(col_widths), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    for i, ratio in enumerate(col_widths):
        table.columns[i].width = Inches(w * ratio)
    header_h = 0.42
    body_h = (h - header_h) / (len(rows) - 1)
    for ri, row in enumerate(rows):
        table.rows[ri].height = Inches(header_h if ri == 0 else body_h)
        for ci, value in enumerate(row):
            set_cell(table.cell(ri, ci), value, header=(ri == 0), shade=(ri % 2 == 0), size=font_size)
    return table_shape


def add_callout(slide, text, x, y, w, h, size=14):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = GOLD_PALE; shape.line.color.rgb = RGBColor(226, 195, 112)
    box = slide.shapes.add_textbox(Inches(x + 0.16), Inches(y + 0.10), Inches(w - 0.32), Inches(h - 0.18))
    box.text_frame.word_wrap = True
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = box.text_frame.paragraphs[0]; p.space_after = Pt(0)
    r = p.add_run(); r.text = text; r.font.name = "Arial"; r.font.size = Pt(size); r.font.color.rgb = DARK


def add_card(slide, x, title, body, *, width=5.65, body_size=15):
    y, h = 1.78, 4.95
    base = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(h))
    base.fill.solid(); base.fill.fore_color.rgb = PALE; base.line.fill.background()
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(0.62))
    band.fill.solid(); band.fill.fore_color.rgb = TEAL; band.line.fill.background()
    titlebox = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.07), Inches(width - 0.36), Inches(0.48))
    p = titlebox.text_frame.paragraphs[0]; p.space_after = Pt(0)
    r = p.add_run(); r.text = title; r.font.name = "Arial"; r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = WHITE
    bodybox = slide.shapes.add_textbox(Inches(x + 0.22), Inches(y + 0.82), Inches(width - 0.44), Inches(h - 1.00))
    tf = bodybox.text_frame; tf.clear(); tf.word_wrap = True
    for i, line in enumerate(body):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + line
        p.font.name = "Arial"; p.font.size = Pt(body_size); p.font.color.rgb = DARK
        p.space_after = Pt(10); p.level = 0


prs = Presentation(str(SOURCE))
if len(prs.slides) != 12:
    raise RuntimeError(f"Expected 12 source slides, found {len(prs.slides)}")

# Slide 7: distinguish repository implementation from a verified reproduction.
slide7 = prs.slides[6]
for shape in slide7.shapes:
    if not shape.has_text_frame:
        continue
    if shape.text.strip() == "Base Paper Implementation":
        shape.text = "4D-GS Baseline Implementation"
    elif "Reproduced the full 4D-GS pipeline" in shape.text:
        shape.text = (
            "Repository implementation based on the 4D-GS codebase\n"
            "• Deformation: HexPlane grid + MLP for time-conditioned motion (deformation.py, hexplane.py).\n"
            "• Gaussian model: position, opacity, covariance and appearance (gaussian_model.py).\n"
            "• Training, rendering, dataset loaders and CUDA submodule integration are present.\n"
            "• Phase-1 notebook documents the intended D-NeRF workflow.\n\n"
            "Evidence boundary: source and workflow are present; present a fresh run with its config, logs and checkpoint before claiming a reproduced benchmark."
        )
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(16)
    elif shape.text.strip().startswith("Result: the baseline runs end to end"):
        shape.text = ""

# Slide 9: replace absolute completion claims with implementation status and caveats.
slide9 = prs.slides[8]
for shape in slide9.shapes:
    if shape.has_text_frame and shape.text.strip() == "Progress vs. Original Proposal":
        shape.text = "Implementation Status vs. Original Proposal"
    elif shape.has_text_frame and shape.text.startswith("What we promised"):
        shape.text = "Repository implementation status; execution and metric validation are tracked separately"
    elif shape.has_table:
        rows = [
            ("Component", "Status", "Evidence / boundary"),
            ("4D-GS training/render path", "In repo", "Dynamic model, training and rendering modules are present."),
            ("SAM 2 mask extraction", "Implemented", "Per-frame extraction path; weights and runtime setup required."),
            ("CLIP features and labels", "Implemented", "OpenCLIP image/text features and prompt-based labeling path."),
            ("Semantic Gaussian supervision", "Implemented", "Loader and training/render hooks exist; fresh metrics not verified here."),
            ("Grounding DINO", "Optional", "Integration path exists; real detector requires model weights/config."),
            ("Gaussian editing", "Prototype", "Class selection and position edits; not a full ARAP energy solver."),
            ("Web demo", "Illustrative", "Not connected to trained dynamic 4D-GS inference."),
        ]
        for ri, row in enumerate(rows):
            for ci, value in enumerate(row):
                set_cell(shape.table.cell(ri, ci), value, header=(ri == 0), shade=(ri % 2 == 0), size=11)
        shape.table.columns[0].width = Inches(3.65)
        shape.table.columns[1].width = Inches(1.45)
        shape.table.columns[2].width = Inches(6.5)
        for ri, row in enumerate(shape.table.rows):
            row.height = Inches(0.44 if ri == 0 else 0.49)
    elif shape.has_text_frame and shape.text.startswith("Summary:"):
        shape.text = "Status means the code path is present—not that every stage has a reproducible end-to-end run. See the next slides for evidence and validation needs."
        shape.top = Inches(5.95); shape.height = Inches(0.65)
        for paragraph in shape.text_frame.paragraphs:
            paragraph.font.size = Pt(15)

# Slide 10: replace unverified model comparison and precise metrics with evidence-level reporting.
slide10 = prs.slides[9]
for shape in list(slide10.shapes):
    if shape.has_table:
        shape._element.getparent().remove(shape._element)
    elif shape.has_text_frame and shape.text.startswith("Final Quantitative Achievements"):
        shape._element.getparent().remove(shape._element)
for shape in slide10.shapes:
    if shape.has_text_frame and shape.text.strip() == "6. Result Analysis":
        shape.text = "Prototype Results & Evaluation Status"
    elif shape.has_text_frame and shape.text.strip() == "Comparison with existing models":
        shape.text = "Implemented functionality is distinct from reproducible quantitative evidence"
add_table(slide10, 0.85, 1.75, 11.6, 2.75, [
    ("Capability", "Repository evidence", "Safe interpretation"),
    ("Dynamic reconstruction", "4D-GS model, training and render code", "Pipeline is present; reproduce with a recorded setup."),
    ("2D semantics", "SAM 2 + OpenCLIP extraction modules", "Extraction path is implemented; outputs depend on weights/prompts."),
    ("Semantic lifting", "Semantic loader and train/render hooks", "End-to-end accuracy needs rendered predictions and ground truth."),
    ("Object editing", "Class selection and Gaussian position transforms", "Prototype behavior; full 4D/covariance consistency needs testing."),
], [0.22, 0.34, 0.44], font_size=12)
add_callout(
    slide10,
    "The checked-in evaluator can report pixel-space label-map IoU when predictions and ground truth are supplied. It does not verify PSNR/SSIM, temporal consistency, FPS or editing latency. Historical draft values for these metrics are unverified; do not present them as final results without archived runs, configs and metric definitions.",
    0.85, 4.78, 11.6, 1.55, size=14,
)

# New slide: explicitly separate implemented code from validated evidence.
new_slide = prs.slides.add_slide(prs.slide_layouts[0])
add_title(new_slide, "What Is Implemented / What Is Validated", "A source-level implementation review is not the same as experimental validation", 11)
add_card(new_slide, 0.85, "Implemented in the repository", [
    "4D-GS-based dynamic reconstruction and rendering code path.",
    "SAM 2 masks and OpenCLIP image/text feature extraction.",
    "Semantic data loading plus training/render integration hooks.",
    "Gaussian selection/editing prototype; separate static multi-view projection and voting pipeline.",
], width=5.65, body_size=15)
add_card(new_slide, 6.83, "Validated / still to validate", [
    "Source modules, command-line entry points and phase notebooks are present.",
    "A fresh end-to-end run is not established by source presence alone.",
    "Current evaluator needs rendered labels and ground truth; it measures pixel-label IoU only.",
    "Re-run with archived configs, checkpoints and logs; separately define and test image quality, temporal consistency, speed and editing metrics.",
], width=5.65, body_size=14)

# Move the added slide to follow the evaluation slide (before the conclusion).
sld_ids = prs.slides._sldIdLst
new_id = sld_ids[-1]
sld_ids.remove(new_id)
sld_ids.insert(10, new_id)

# Slide 12: revise the existing conclusion/limitations cards for honest, concise claims.
slide12 = prs.slides[11]
for shape in slide12.shapes:
    if not shape.has_text_frame:
        continue
    text = shape.text.strip()
    if text == "Completed Contributions, Limitations & Future Directions":
        shape.text = "Summary, Limitations & Next Steps"
    elif text == "Comprehensive summary of deliverables, innovations, and impact":
        shape.text = "A promising research prototype; quantitative claims require reproducible evaluation"
    elif text == "Completed Deliverables":
        shape.text = "Implemented Prototype"
    elif text == "Conclusion":
        shape.text = "Takeaway"
    elif text.startswith("• Full 4D-GS baseline"):
        tf = shape.text_frame; tf.clear(); tf.word_wrap = True
        items = [
            "4D-GS-based dynamic reconstruction code path.",
            "SAM 2 + CLIP semantic extraction and supervision path.",
            "Semantic loading/rendering integration.",
            "Gaussian selection and editing prototype.",
        ]
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = "• " + item; p.font.name = "Arial"; p.font.size = Pt(14); p.space_after = Pt(7)
    elif text.startswith("• Storage:"):
        tf = shape.text_frame; tf.clear(); tf.word_wrap = True
        items = [
            "Semantic quality depends on 2D model masks, prompts and occlusion.",
            "Fresh training/evaluation runs and ground-truth labels are needed.",
            "Editing changes Gaussian positions; full temporal/covariance consistency is unverified.",
            "Static multi-view voting is a separate pipeline from dynamic 4D training.",
        ]
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = "• " + item; p.font.name = "Arial"; p.font.size = Pt(13); p.space_after = Pt(6)
    elif text.startswith("• Developed the first"):
        tf = shape.text_frame; tf.clear(); tf.word_wrap = True
        items = [
            "This work is a prototype for semantic-aware dynamic Gaussian reconstruction and editing.",
            "Its central idea is to connect 2D semantic supervision to a dynamic Gaussian workflow.",
            "Next: run reproducible benchmarks against the unmodified 4D-GS baseline under matched settings.",
        ]
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = "• " + item; p.font.name = "Arial"; p.font.size = Pt(14); p.space_after = Pt(8)

# Renumber only the footer page-number boxes after inserting slide 11.
for page, slide in enumerate(prs.slides, start=1):
    for shape in slide.shapes:
        if shape.has_text_frame and shape.top >= Inches(6.8) and shape.left >= Inches(11.5):
            if shape.text.strip().isdigit():
                shape.text = str(page)

prs.save(str(OUTPUT))
print(f"Saved revised deck: {OUTPUT}")
print(f"Slides: {len(prs.slides)}")
